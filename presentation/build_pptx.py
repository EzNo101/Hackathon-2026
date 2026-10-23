from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"

BLUE = RGBColor(0x1F, 0x4E, 0x79)
GREEN = RGBColor(0x2A, 0x9D, 0x8F)
ORANGE = RGBColor(0xE7, 0x6F, 0x51)
INK = RGBColor(0x1B, 0x2A, 0x38)
GRAY = RGBColor(0x5B, 0x6B, 0x7B)
LIGHT = RGBColor(0xEA, 0xF1, 0xF8)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_text(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    size: int = 18,
    color: RGBColor = INK,
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    anchor: MSO_ANCHOR = MSO_ANCHOR.TOP,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = align
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.font.name = FONT


def add_bullets(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    items: list[str],
    size: int = 18,
    color: RGBColor = INK,
    space_after: int = 7,
) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.font.name = FONT
        p.space_after = Pt(space_after)


def add_title(slide, text: str) -> None:
    add_text(slide, 0.7, 0.38, 11.9, 0.8, text, size=30, color=BLUE, bold=True)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.18), Inches(1.1), Inches(0.05)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = GREEN
    bar.line.fill.background()


def add_footer(slide, number: int) -> None:
    add_text(
        slide,
        0.7,
        7.05,
        8.0,
        0.3,
        "COMPASS · Hackathon 2026 · Děčín",
        size=10,
        color=GRAY,
    )
    add_text(
        slide,
        12.0,
        7.05,
        0.7,
        0.3,
        str(number),
        size=10,
        color=GRAY,
        align=PP_ALIGN.RIGHT,
    )


def add_table(
    slide,
    data: list[list[str]],
    left: float,
    top: float,
    width: float,
    font_size: int = 15,
    row_height: float = 0.36,
    col_widths: list[float] | None = None,
) -> None:
    rows, cols = len(data), len(data[0])
    shape = slide.shapes.add_table(
        rows, cols, Inches(left), Inches(top), Inches(width), Inches(row_height * rows)
    )
    table = shape.table
    if col_widths:
        for i, w in enumerate(col_widths):
            table.columns[i].width = Inches(w)
    for r in range(rows):
        table.rows[r].height = Inches(row_height)
        for c in range(cols):
            cell = table.cell(r, c)
            cell.text = str(data[r][c])
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.margin_top = Inches(0.02)
            cell.margin_bottom = Inches(0.02)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = cell.text_frame.paragraphs[0]
            p.font.size = Pt(font_size)
            p.font.name = FONT
            p.font.bold = r == 0
            p.font.color.rgb = WHITE if r == 0 else INK
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (
                LIGHT if r % 2 == 0 else WHITE
            )


def add_callout(
    slide,
    left: float,
    top: float,
    width: float,
    height: float,
    text: str,
    color: RGBColor = GREEN,
    size: int = 18,
) -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.text = text
    p.alignment = PP_ALIGN.CENTER
    p.font.size = Pt(size)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.font.name = FONT


def add_picture(slide, name: str, left: float, top: float, width: float) -> None:
    slide.shapes.add_picture(str(IMG / name), Inches(left), Inches(top), width=Inches(width))


def content_slide(prs, title: str, number: int):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_title(slide, title)
    add_footer(slide, number)
    return slide


def build() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 — title
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = BLUE
    bg.line.fill.background()
    add_text(
        slide,
        0.9,
        2.0,
        11.6,
        1.6,
        "Rekonstrukce tvaru plazmatu\nna tokamaku COMPASS",
        size=44,
        color=WHITE,
        bold=True,
    )
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.95), Inches(3.95), Inches(2.2), Inches(0.07))
    line.fill.solid()
    line.fill.fore_color.rgb = GREEN
    line.line.fill.background()
    add_text(
        slide,
        0.95,
        4.25,
        11.0,
        0.6,
        "Z magnetických diagnostik (a volitelně z kamery)",
        size=24,
        color=RGBColor(0xD8, 0xE6, 0xF3),
    )
    add_text(slide, 0.95, 6.35, 8.0, 0.4, "Hackathon 2026 · Děčín", size=16, color=RGBColor(0xD8, 0xE6, 0xF3))

    # 2 — goal
    slide = content_slide(prs, "Cíl", 2)
    add_bullets(
        slide,
        0.7,
        1.5,
        11.9,
        2.6,
        [
            "Z 24 magnetických signálů (IPR01–16, FL01–08) předpovědět:",
            "•  10 parametrů tvaru a polohy plazmatu",
            "•  X-point a 2 strike points (jen diverted plazma)",
            "•  bod kontaktu s limiterem (R, Z) (jen limited plazma)",
            "•  binární režim is_limited",
        ],
        size=19,
    )
    add_table(
        slide,
        [
            ["Data", "Rozsah"],
            ["equil_data.csv", "42 972 řádků · 149 výbojů · 57 sloupců"],
            ["camera_data.csv", "116 142 snímků · 78 výbojů · 640×500 RGB"],
        ],
        0.7,
        4.35,
        11.9,
        font_size=17,
        row_height=0.48,
        col_widths=[3.3, 8.6],
    )
    add_callout(
        slide,
        0.7,
        6.05,
        11.9,
        0.55,
        "24 příznaků  →  18 cílových veličin  +  binární režim",
        color=GREEN,
        size=18,
    )

    # 3 — physics
    slide = content_slide(prs, "Fyzika dat: limited vs diverted", 3)
    add_table(
        slide,
        [
            ["Režim", "Co plazma dělá", "Vyplněné cíle"],
            ["limited (is_limited=True)", "dotýká se limiteru", "R_limiter, Z_limiter"],
            ["diverted (is_limited=False)", "má X-point, jde na divertor", "X-point + strike points"],
        ],
        0.7,
        1.5,
        11.9,
        font_size=17,
        row_height=0.52,
        col_widths=[3.9, 4.2, 3.8],
    )
    add_bullets(
        slide,
        0.7,
        3.35,
        11.9,
        3.2,
        [
            "Chybějící hodnoty nejsou chyba dat — jsou fyzikální.",
            "Limited: 19 012 řádků · diverted: 23 960 řádků.",
            "Strike points = místa, kde magnetické čáry z X-pointu dopadají na stěnu.",
            "R_strike_point_1 < R_strike_point_2 (ověřeno na 100 % řádků).",
        ],
        size=19,
    )

    # 4 — limiter point
    slide = content_slide(prs, "Limiter: jeden bod, dvě souřadnice", 4)
    add_bullets(
        slide,
        0.7,
        1.55,
        6.4,
        3.4,
        [
            "R_limiter = radiální souřadnice",
            "Z_limiter = vertikální souřadnice",
            "Není to „levý a pravý limiter“ — je to jedna kontaktní pointa (R, Z).",
            "73,4 % hodnot R_limiter je 0,347 m → vnitřní prstenec.",
            "Zbytek tvoří jiné body stěny (R až 0,771 m).",
        ],
        size=19,
    )
    add_picture(slide, "limiter_points.png", 7.35, 1.45, 5.2)
    add_callout(
        slide,
        0.7,
        5.75,
        6.4,
        0.7,
        "NaN = plazma se limiteru vůbec nedotýká",
        color=ORANGE,
        size=17,
    )

    # 5 — methodology
    slide = content_slide(prs, "Metodika", 5)
    add_text(slide, 0.7, 1.5, 5.8, 0.4, "Rozdělení dat", size=21, color=BLUE, bold=True)
    add_bullets(
        slide,
        0.7,
        2.0,
        5.8,
        3.4,
        [
            "Podle výbojů, ne po řádcích",
            "119 výbojů train / 30 výbojů test",
            "33 902 / 9 070 řádků",
            "time a discharge nejsou příznaky",
            "Kanonikalizace strike points podle R",
        ],
        size=18,
    )
    add_text(slide, 7.0, 1.5, 5.6, 0.4, "Model", size=21, color=BLUE, bold=True)
    add_bullets(
        slide,
        7.0,
        2.0,
        5.6,
        3.4,
        [
            "StandardScaler + MLP 128×128",
            "Standardizace cílových veličin",
            "GroupKFold(5) podle výbojů",
            "Pevné rozdělení pro finální metriky",
            "Stejná architektura pro všechny skupiny",
        ],
        size=18,
    )
    add_callout(
        slide,
        0.7,
        5.75,
        11.9,
        0.7,
        "Bez úniku dat: žádný výboj není současně v train i test",
        color=BLUE,
        size=18,
    )

    # 6 — shape results
    slide = content_slide(prs, "Výsledky: tvar a poloha", 6)
    add_table(
        slide,
        [
            ["Cíl", "MAE", "R²"],
            ["R_mag_axis", "0,34 mm", "0,985"],
            ["Z_mag_axis", "0,43 mm", "0,995"],
            ["elongation", "0,0058", "0,998"],
            ["a", "0,50 mm", "0,992"],
            ["plasma_current", "1 021 A", "0,9996"],
        ],
        0.7,
        1.5,
        6.1,
        font_size=16,
        row_height=0.38,
        col_widths=[2.6, 1.7, 1.8],
    )
    add_text(
        slide,
        7.2,
        1.55,
        5.4,
        1.2,
        "GroupKFold(5):\nR 0,61 ± 0,19 mm\nZ 0,70 ± 0,16 mm",
        size=18,
        color=GRAY,
    )
    add_picture(slide, "parity.png", 2.2, 3.95, 9.0)

    # 7 — xpoint
    slide = content_slide(prs, "X-point a strike points", 7)
    add_table(
        slide,
        [
            ["Cíl", "MAE", "R²"],
            ["R_x_point", "0,24 mm", "0,9965"],
            ["Z_x_point", "0,27 mm", "0,9972"],
            ["R_strike_point_1", "0,30 mm", "0,9980"],
            ["Z_strike_point_1", "0,13 mm", "0,9972"],
            ["R_strike_point_2", "0,25 mm", "0,9953"],
            ["Z_strike_point_2", "0,06 mm", "0,9886"],
        ],
        0.7,
        1.5,
        8.2,
        font_size=17,
        row_height=0.44,
        col_widths=[3.8, 2.1, 2.3],
    )
    add_callout(
        slide,
        9.25,
        2.1,
        3.3,
        1.5,
        "R² > 0,988\npro všech 6 cílů",
        color=GREEN,
        size=21,
    )
    add_text(
        slide,
        9.25,
        3.9,
        3.3,
        1.4,
        "5 438 diverted řádků\ntestovacích výbojů",
        size=16,
        color=GRAY,
    )

    # 8 — is_limited
    slide = content_slide(prs, "Klasifikace režimu is_limited", 8)
    add_text(
        slide,
        0.7,
        1.5,
        11.9,
        0.6,
        "MLP klasifikátor na stejných 24 magnetických příznacích",
        size=19,
        color=GRAY,
    )
    for i, (value, label) in enumerate(
        [("99,1 %", "accuracy"), ("0,9996", "ROC-AUC"), ("0,989", "F1")]
    ):
        add_callout(
            slide,
            0.7 + i * 4.0,
            2.45,
            3.6,
            1.35,
            f"{value}\n{label}",
            color=GREEN if i == 0 else BLUE,
            size=24,
        )
    add_bullets(
        slide,
        0.7,
        4.35,
        11.9,
        1.8,
        [
            "Balanced accuracy: 99,1 %",
            "Režim je z magnetických signálů prakticky dokonale rozlišitelný.",
            "Dává smysl použít ho jako první stupeň před regresí limiteru.",
        ],
        size=19,
    )

    # 9 — limiter two-stage
    slide = content_slide(prs, "Limiter: dvoustupňový model", 9)
    add_bullets(
        slide,
        0.7,
        1.5,
        6.3,
        2.0,
        [
            "1. Klasifikátor: inner (R ≤ 0,35) vs other.",
            "2. Inner: R = 0,347 (konstanta) + model pro Z.",
            "3. Other: samostatné modely pro R a Z.",
        ],
        size=18,
    )
    add_table(
        slide,
        [
            ["Cíl", "Baseline", "Two-stage", "Změna"],
            ["R_limiter", "16,2 mm", "8,0 mm", "−50,4 %"],
            ["Z_limiter", "22,8 mm", "15,0 mm", "−34,2 %"],
        ],
        0.7,
        3.65,
        6.3,
        font_size=16,
        row_height=0.48,
        col_widths=[1.8, 1.5, 1.5, 1.5],
    )
    add_picture(slide, "limiter_cv.png", 7.3, 1.7, 5.3)
    add_callout(
        slide,
        0.7,
        5.75,
        6.3,
        0.7,
        "GroupKFold(5), MAE v mm",
        color=BLUE,
        size=17,
    )

    # 10 — mode
    slide = content_slide(prs, "Negativní výsledek: mode jako příznak", 10)
    add_text(
        slide,
        0.7,
        1.5,
        11.9,
        0.8,
        "Přidání one-hot mode (L/H/ELM) z kamery k magnetickým příznakům:",
        size=19,
        color=GRAY,
    )
    add_table(
        slide,
        [
            ["Vyhodnocení", "Změna průměrného MAE"],
            ["Jedno pevné rozdělení", "−5,2 % (zdánlivý zisk)"],
            ["GroupKFold(5)", "+4,2 % (zhoršení)"],
        ],
        0.7,
        2.5,
        11.9,
        font_size=18,
        row_height=0.55,
        col_widths=[6.0, 5.9],
    )
    add_bullets(
        slide,
        0.7,
        4.6,
        11.9,
        1.3,
        [
            "Signál je nestabilní → model se přeučuje na konkrétní výboje.",
            "mode není ve finálním modelu.",
        ],
        size=19,
    )
    add_callout(
        slide,
        0.7,
        6.0,
        11.9,
        0.6,
        "Kamera zůstává jen pro analýzu režimu, ne jako příznak",
        color=ORANGE,
        size=17,
    )

    # 11 — speed and limits
    slide = content_slide(prs, "Rychlost a omezení", 11)
    add_text(slide, 0.7, 1.5, 5.8, 0.4, "Rychlost", size=21, color=GREEN, bold=True)
    add_bullets(
        slide,
        0.7,
        2.0,
        5.8,
        2.8,
        [
            "Trénink všech modelů: ~13 s",
            "Infernce všech modelů: ~1,5 ms / vzorek",
            "Vhodné pro real-time řízení",
        ],
        size=18,
    )
    add_text(slide, 7.0, 1.5, 5.6, 0.4, "Omezení", size=21, color=ORANGE, bold=True)
    add_bullets(
        slide,
        7.0,
        2.0,
        5.6,
        3.2,
        [
            "Jedna architektura (MLP), bez ladění hyperparametrů",
            "Jen magnetické příznaky",
            "R_limiter R² ~0,68 je důsledek konstantní predikce",
            "Testováno na 30 výbojích",
        ],
        size=18,
    )
    add_callout(
        slide,
        0.7,
        5.75,
        11.9,
        0.7,
        "~1,5 ms na vzorek  →  použitelné v reálném čase",
        color=GREEN,
        size=19,
    )

    # 12 — conclusions
    slide = content_slide(prs, "Závěry", 12)
    add_bullets(
        slide,
        0.7,
        1.5,
        7.2,
        4.6,
        [
            "Tvar a poloha: R² 0,985–0,9996, chyba centra < 0,5 mm",
            "X-point / strike points: R² > 0,988",
            "Režim: accuracy 99,1 %, ROC-AUC 0,9996",
            "Limiter: dvoustupňový přístup snížil MAE o 34–50 %",
            "Latence ~1,5 ms → použitelné v reálném čase",
            "mode jako příznak se neosvědčil",
        ],
        size=18,
    )
    add_text(slide, 8.3, 1.5, 4.3, 0.4, "Další kroky", size=21, color=BLUE, bold=True)
    add_bullets(
        slide,
        8.3,
        2.05,
        4.3,
        3.6,
        [
            "1. Odolnost proti šumu",
            "2. Camera CNN",
            "3. Multimodální fúze",
            "4. Více výbojů pro validaci",
        ],
        size=17,
    )
    add_callout(
        slide,
        8.3,
        5.75,
        4.3,
        0.8,
        "Sešity + tato prezentace\njsou hotové",
        color=GREEN,
        size=16,
    )

    out = HERE / "prezentace.pptx"
    prs.save(out)
    print(f"saved: {out}")


if __name__ == "__main__":
    build()
